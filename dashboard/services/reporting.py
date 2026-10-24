"""Reports: one generic structure, several renderings.

A report is {"type", "title", "scope", "period", "subtitle", "model", "hero", "sections", "footer", "state", "filename"}.
Sections are {"heading", "columns", "rows", "row_fmts"?, "row_styles"?, "note"?, "collapsed"?}; columns are {"label", "fmt"} and
rows are lists of *raw* values (numbers stay numbers). build_* functions produce that structure from kpis.py / the ledger -- no
number is computed here that isn't derived on read from the same functions every page uses -- and to_csv/to_xlsx/to_pdf/
format_cell render it, so the on-screen preview, the spreadsheet and the PDF can never disagree with each other.

Vocabulary (the same everywhere, UI / PDF / CSV / Excel):
  Company Revenue          operated-property revenue + Management Fee Earned on managed properties
  Operated property costs     recorded costs of operated properties
  Property Profit             Company Revenue - Operated property costs        (Business Costs NOT included)
  Business Costs              recorded costs of the shared overhead cost centre
  Company Net Profit       Property Profit - Business Costs
  Gross Booking Revenue       what guests paid; on a managed property it mostly belongs to the owner (owner-side)
  Owner-side balance          Gross Booking Revenue - owner-side costs - the management fee: owner economics, never company profit"""
import csv
import datetime
import io
import re

import services.kpis as kpis
from services import completeness
from services.common import MONTH_NAMES, get_properties, get_property, is_managed, pct_delta, short_name

REPORT_TYPES = {
    "portfolio_monthly": "Portfolio monthly",
    "annual_portfolio": "Portfolio annual",
    "property_monthly": "Property statement",
    "expense": "Expense report",
}
# one line each: what the report contains
REPORT_CONTENTS = {
    "portfolio_monthly": "Company Revenue to Company Net Profit for one month, operations, recorded expenses by cost centre, and each property.",
    "annual_portfolio": "The same statement for a year, month by month and against last year, with each property.",
    "property_monthly": "One property for one month. Operated: Revenue to Property Profit. Managed: owner-side statement and Management Fee Earned.",
    "expense": "Every recorded expense over a range: totals, cost centres, categories, vendors and the ledger.",
}
TYPE_SLUG = {"portfolio_monthly": "Portfolio-Monthly", "annual_portfolio": "Portfolio-Annual", "property_monthly": "Property-Statement", "expense": "Expense-Report"}
MINUS = "\u2212"


# ---------------------------------------------------------------- formatting

def _money(value, places):
    r = round(value, places)
    return (MINUS if r < 0 else "") + f"\u00a3{abs(r):,.{places}f}"


def format_cell(value, fmt):
    if value is None or value == "":
        return "-"
    if fmt == "money":
        return _money(value, 0)
    if fmt == "money2":
        return _money(value, 2)
    if fmt == "pct":
        return f"{value:.0f}%"
    if fmt == "delta":
        return f"{value:+.1f}%".replace("-", MINUS)
    if fmt == "int":
        return f"{value:,.0f}"
    if fmt == "num1":
        return f"{value:,.1f}"
    return str(value)


def _col(label, fmt="text"):
    return {"label": label, "fmt": fmt}


def _pretty(category):
    return (category or "").replace("_", " ").capitalize()


def _slug(text):
    return "-".join(re.findall(r"[A-Za-z0-9]+", text))


def _today():
    d = datetime.date.today()
    return f"{d.day} {MONTH_NAMES[d.month][:3]} {d.year}"


def _source_text(conn):
    """Where the figures come from, quietly: the recorded dashboard data and the latest applied workbook import (if any)."""
    try:
        row = conn.execute("SELECT period FROM import_batches WHERE status='applied' AND kind='workbook' AND period IS NOT NULL ORDER BY id DESC LIMIT 1").fetchone()
    except Exception:                       # a demo database that predates the import tables
        row = None
    text = "Recorded dashboard data"
    if row and re.match(r"^\d{4}-\d{2}", row["period"] or ""):
        text += f" \u00b7 latest workbook import {MONTH_NAMES[int(row['period'][5:7])]} {row['period'][:4]}"
    return text


def _footer(conn, scope, period):
    return [["Generated", _today()], ["Period", period], ["Scope", scope], ["Source", _source_text(conn)]]


# ------------------------------------------------------------------ helpers over the existing KPI functions

def _company(conn, start, end):
    """Business Costs and Company Net Profit from the Overview's own function (Revenue - Operated Property Costs - Business Costs)."""
    from routes.overview import company_net_profit      # the established definition lives with the Overview; imported lazily
    return company_net_profit(conn, start, end)


def _portfolio_values(conn, start, end):
    adj = kpis.adjusted_kpi_snapshot(conn, None, start, end)
    co = _company(conn, start, end)
    return {"revenue": adj["revenue"], "operated": adj["costs"], "pp": adj["net_profit"], "margin": adj["margin"] * 100,
            "business": co["business_costs"], "net": co["net_profit"],
            "occupancy": adj["occupancy"] * 100, "adr": adj["adr"], "revpar": adj["revpar"], "nights": adj["booked_nights"]}


def _fee_recorded(conn, pid, start, end):
    return conn.execute("""SELECT COALESCE(SUM(amount),0) FROM transactions
                           WHERE property_id=? AND direction='expense' AND category='management_fee' AND date>=? AND date<?""", (pid, start, end)).fetchone()[0]


def _property_values(conn, pid, start, end, managed):
    snap = kpis.kpi_snapshot(conn, pid, start, end)
    v = {"revenue": snap["revenue"], "costs": snap["costs"], "profit": snap["net_profit"], "margin": snap["margin"] * 100,
         "occupancy": snap["occupancy"] * 100, "adr": snap["adr"], "revpar": snap["revpar"], "nights": snap["booked_nights"]}
    if managed:
        fee = _fee_recorded(conn, pid, start, end)
        v.update(fee_recorded=fee, owner_costs=snap["costs"] - fee, balance=snap["net_profit"], fee_earned=kpis.business_income(conn, pid, start, end))
    return v


def _row(label, key, cur, prev, last, fmt, min_base=1):
    """[label, value, vs prior, vs last year] plus the format of the value."""
    return [label, cur[key], pct_delta(cur[key], prev[key], min_base=min_base), pct_delta(cur[key], last[key], min_base=min_base)], fmt


def _statement(heading, lines, cur, prev, last, note=None):
    """lines: [(label, key, fmt, min_base, style)] -> a section whose value column carries a format per row."""
    rows, fmts, styles = [], [], []
    for label, key, fmt, min_base, style in lines:
        row, f = _row(label, key, cur, prev, last, fmt, min_base)
        rows.append(row); fmts.append(f); styles.append(style)
    return {"heading": heading, "columns": [_col("Item"), _col("This period", "mixed"), _col("vs prior period", "delta"), _col("vs same period last year", "delta")],
            "rows": rows, "row_fmts": fmts, "row_styles": styles, "note": note}


OPS = [("Occupancy", "occupancy", "pct", 5, ""), ("ADR", "adr", "money", 20, ""), ("RevPAR", "revpar", "money", 20, ""), ("Booked nights", "nights", "int", 2, "")]
DELTA_NOTE = "Percentage changes are omitted when the comparison period's figure is too small to compare meaningfully."


def _category_table(conn, start, end, property_id=None, heading="Expenses by category", note=None):
    clause, params = ("AND property_id=?", (property_id,)) if property_id else ("", ())
    data = conn.execute(
        f"""SELECT category, SUM(amount) amt, COUNT(*) n FROM transactions
            WHERE direction='expense' AND category != 'reconciliation' AND date>=? AND date<? {clause}
            GROUP BY category ORDER BY amt DESC""",
        (start, end, *params),
    ).fetchall()
    total = sum(r["amt"] for r in data) or 0
    rows = [[_pretty(r["category"]), r["n"], r["amt"], (r["amt"] / total * 100) if total else None] for r in data]
    return {"heading": heading,
            "columns": [_col("Category"), _col("Items", "int"), _col("Amount", "money"), _col("Share", "pct")],
            "rows": rows, "note": note}


def _ledger_table(conn, start, end, property_id=None, limit=1000, heading="Transactions"):
    clause, params = ("AND t.property_id=?", (property_id,)) if property_id else ("", ())
    data = conn.execute(
        f"""SELECT t.date, t.vendor, t.description, t.category, t.capex, t.direction, t.amount, p.name pname
            FROM transactions t JOIN properties p ON p.id = t.property_id
            WHERE t.date>=? AND t.date<? {clause} ORDER BY t.date, t.id LIMIT ?""",
        (start, end, *params, limit),
    ).fetchall()
    rows = []
    for r in data:
        kind = "Income" if r["direction"] == "income" else ("Capex" if r["capex"] else "Opex")
        cat = _pretty(r["category"]) + (" (historical adjustment)" if r["category"] == "reconciliation" else "")
        rows.append([r["date"], r["vendor"] or "", short_name(r["pname"]), cat, kind, r["amount"]])
    return {"heading": heading, "collapsed": True,
            "columns": [_col("Date"), _col("Vendor"), _col("Property"), _col("Category"), _col("Type"), _col("Amount", "money2")],
            "rows": rows,
            "note": f"Showing up to {limit} transactions." if len(rows) >= limit else None}


def _vendor_table(conn, start, end, property_id=None, heading="Top vendors"):
    clause, params = ("AND t.property_id=?", (property_id,)) if property_id else ("", ())
    data = conn.execute(
        f"""SELECT v.name, SUM(t.amount) amt, COUNT(*) n FROM transactions t JOIN vendors v ON v.id = t.vendor_id
            WHERE t.direction='expense' AND t.category != 'reconciliation' AND t.date>=? AND t.date<? {clause}
            GROUP BY v.id ORDER BY amt DESC LIMIT 15""",
        (start, end, *params),
    ).fetchall()
    return {"heading": heading, "columns": [_col("Vendor"), _col("Items", "int"), _col("Amount", "money")],
            "rows": [[r["name"], r["n"], r["amt"]] for r in data]}


def _cost_centres(conn, start, end, property_id=None, managed=None):
    """Recorded expenses split by whose they are, using the Expenses page's own definitions plus the management fee on its own line.
    Returned only if the parts add up to the total recorded expenses exactly -- never a plug."""
    from routes.expenses import _costs                    # the Expenses page's definitions; imported lazily
    total = kpis.costs(conn, property_id, start, end)
    fee_clause, fee_params = ("AND property_id=?", (property_id,)) if property_id else ("", ())
    fees = conn.execute(f"SELECT COALESCE(SUM(amount),0) FROM transactions WHERE direction='expense' AND category='management_fee' AND date>=? AND date<? {fee_clause}",
                        (start, end, *fee_params)).fetchone()[0]
    if property_id:
        rows = [["Owner-side costs (excluding the management fee)" if managed else "Property costs", _costs(conn, start, end, property_id=property_id)],
                ["Management fee recorded", fees]]
    else:
        rows = [["Operated property costs", _costs(conn, start, end, scope="property", model="operated")],
                ["Managed properties' owner-side costs (not company costs)", _costs(conn, start, end, scope="property", model="managed")],
                ["Management fees recorded (The company's own income)", fees],
                ["Business Costs", _costs(conn, start, end, scope="business")]]
    if abs(sum(r[1] for r in rows) - total) > 0.005:
        return None
    rows.append(["Total recorded expenses", total])
    return {"heading": "Recorded expenses by cost centre", "columns": [_col("Cost centre"), _col("Amount", "money")], "rows": rows,
            "row_styles": [""] * (len(rows) - 1) + ["total"],
            "note": "Only Operated property costs and Business Costs enter the company's profit figures. Owner-side costs belong to the owners; the management fee is company revenue."}


def _state_report(base, kind, text):
    base.update(state={"kind": kind, "text": text}, hero=[], sections=[])
    return base


def _base(rtype, scope, period, conn, model=None, filename_scope=None, filename_period=None):
    return {"type": rtype, "title": {"portfolio_monthly": "Portfolio monthly report", "annual_portfolio": "Annual portfolio report",
                                     "property_monthly": "Property statement", "expense": "Expense report"}[rtype],
            "scope": scope, "period": period, "subtitle": f"{scope} \u00b7 {period}", "model": model, "hero": [], "sections": [], "state": None,
            "footer": _footer(conn, scope, period),
            "filename": f"Fernhill_{_slug(filename_scope or scope)}_{filename_period}_{TYPE_SLUG[rtype]}"}


# ------------------------------------------------------------------- reports

def _month_ranges(year, month):
    start, end = kpis.month_bounds(year, month)
    p = kpis.prior_month(year, month)
    ly = (year - 1, month)
    return (start, end, *kpis.month_bounds(*p), *kpis.month_bounds(*ly))


PORTFOLIO_LINES = [("Company Revenue", "revenue", "money", 100, ""), ("Operated property costs", "operated", "money", 100, ""),
                   ("Property Profit", "pp", "money", 1000, "sub"), ("Business Costs", "business", "money", 100, ""),
                   ("Company Net Profit", "net", "money", 1000, "total"), ("Property Profit margin", "margin", "pct", 5, "")]
PORTFOLIO_NOTE = ("Company Revenue is operated-property revenue plus Management Fee Earned (the recorded fee, or the stored fee % of Gross Booking Revenue where none is recorded). "
                  "Managed owners' booking revenue and costs are not included. Property Profit does not include Business Costs; Company Net Profit does. "
                  "Based on recorded dashboard data. " + DELTA_NOTE)


def _by_property(conn, start, end, heading="By property"):
    rows = []
    for p in get_properties(conn, include_overhead=False, period=(start, end)):
        s = kpis.adjusted_kpi_snapshot(conn, p["id"], start, end)
        fee = p["management_fee_pct"]
        model = (f"Managed \u00b7 {fee:g}%" if fee else "Managed \u00b7 fee % not set") if is_managed(p) else "Operated"
        rows.append([short_name(p["name"]), s["revenue"], s["costs"], s["net_profit"], s["margin"] * 100, s["occupancy"] * 100, s["adr"], model])
    rows.sort(key=lambda r: r[1], reverse=True)
    return {"heading": heading, "columns": [_col("Property"), _col("Company Revenue", "money"), _col("Operated property costs", "money"), _col("Property Profit", "money"),
                                             _col("Property Profit margin", "pct"), _col("Occupancy", "pct"), _col("ADR", "money"), _col("Model")], "rows": rows,
            "note": "For a managed property, Company Revenue is its Management Fee Earned; its Gross Booking Revenue and owner-side costs are not included."}


def build_portfolio_monthly(conn, year, month):
    s, e, ps, pe, ls, le = _month_ranges(year, month)
    period = f"{MONTH_NAMES[month]} {year}"
    rep = _base("portfolio_monthly", "All properties", period, conn, filename_period=f"{year}-{month:02d}")
    if not completeness.period_has_data(conn, None, s, e):
        return _state_report(rep, "no_data", "No data has been recorded for this period.")
    cur, prev, last = _portfolio_values(conn, s, e), _portfolio_values(conn, ps, pe), _portfolio_values(conn, ls, le)
    rep["hero"] = [{"label": "Company Net Profit", "value": cur["net"], "fmt": "money"}, {"label": "Property Profit", "value": cur["pp"], "fmt": "money"}]
    rep["sections"] = [_statement("Financial summary", PORTFOLIO_LINES, cur, prev, last, PORTFOLIO_NOTE),
                       _statement("Operations", OPS, cur, prev, last), _cost_centres(conn, s, e), _by_property(conn, s, e)]
    rep["sections"] = [x for x in rep["sections"] if x]
    return rep


def build_property_monthly(conn, property_id, year, month):
    prop = get_property(conn, property_id)
    s, e, ps, pe, ls, le = _month_ranges(year, month)
    managed = is_managed(prop)
    period = f"{MONTH_NAMES[month]} {year}"
    rep = _base("property_monthly", short_name(prop["name"]), period, conn, model="managed" if managed else "operated", filename_period=f"{year}-{month:02d}")
    rep["full_name"] = prop["name"]
    why = completeness.not_active(conn, property_id, s, e)
    if why and why["kind"] == "not_started":
        return _state_report(rep, "not_started", f"Not active in this period \u00b7 starts {why['starts']}.")
    if not completeness.period_has_data(conn, property_id, s, e):
        return _state_report(rep, "no_data", "No data has been recorded for this property in this period.")
    cur, prev, last = (_property_values(conn, property_id, a, b, managed) for a, b in ((s, e), (ps, pe), (ls, le)))
    if managed:
        fee = prop["management_fee_pct"]
        estimated = cur["fee_recorded"] == 0 and cur["fee_earned"] > 0
        rep["hero"] = [{"label": "Gross Booking Revenue", "value": cur["revenue"], "fmt": "money", "note": "owner-side"},
                       {"label": "Management Fee Earned", "value": cur["fee_earned"], "fmt": "money", "note": "The company" + (" \u00b7 estimated" if estimated else "")}]
        owner = _statement("Owner-side statement", [("Gross Booking Revenue (owner-side)", "revenue", "money", 100, ""),
                                                    ("Owner-side costs (excluding the management fee)", "owner_costs", "money", 100, ""),
                                                    ("Management fee deducted (recorded)", "fee_recorded", "money", 100, ""),
                                                    ("Owner-side balance", "balance", "money", 1000, "total")], cur, prev, last,
                           "Owner-side balance is Gross Booking Revenue minus recorded owner-side costs and the management fee. It is owner economics: not company profit, not Property Profit, and not an owner payout.")
        co_note = (f"No management fee is recorded for this period, so Management Fee Earned is estimated at {fee:g}% of Gross Booking Revenue." if estimated and fee else
                   ("Recorded management fee." if cur["fee_recorded"] else None))
        company_stmt = _statement("The company", [("Management Fee Earned", "fee_earned", "money", 100, "total")], cur, prev, last, co_note)
        ops = _statement("Operations", OPS, cur, prev, last, "ADR and RevPAR are guest rates, based on Gross Booking Revenue. " + DELTA_NOTE)
        cats = _category_table(conn, s, e, property_id, "Recorded expenses by category", "Owner-side costs, including the management fee recorded against this property.")
        rep["sections"] = [owner, company_stmt, ops, cats, _ledger_table(conn, s, e, property_id, heading="Transactions")]
    else:
        rep["hero"] = [{"label": "Property Profit", "value": cur["profit"], "fmt": "money"}, {"label": "Revenue", "value": cur["revenue"], "fmt": "money"}]
        fin = _statement("Financial summary", [("Revenue", "revenue", "money", 100, ""), ("Property Costs", "costs", "money", 100, ""),
                                               ("Property Profit", "profit", "money", 1000, "total"), ("Property Profit margin", "margin", "pct", 5, "")], cur, prev, last,
                         "Property Profit is Revenue minus this property's own costs. Business Costs are not allocated to a property. " + DELTA_NOTE)
        rep["sections"] = [fin, _statement("Operations", OPS, cur, prev, last), _category_table(conn, s, e, property_id, "Property costs by category"),
                           _ledger_table(conn, s, e, property_id, heading="Transactions")]
    return rep


def build_annual_portfolio(conn, year):
    s, e = kpis.range_bounds(year, 1, year, 12)
    ls, le = kpis.range_bounds(year - 1, 1, year - 1, 12)
    rep = _base("annual_portfolio", "All properties", str(year), conn, filename_period=str(year))
    if not completeness.period_has_data(conn, None, s, e):
        return _state_report(rep, "no_data", "No data has been recorded for this year.")
    monthly = []
    months = kpis.months_with_data(conn, None)
    for m in range(1, 13):
        if f"{year}-{m:02d}" not in months:
            continue
        ms, me = kpis.month_bounds(year, m)
        v = _portfolio_values(conn, ms, me)
        monthly.append([MONTH_NAMES[m], v["revenue"], v["operated"], v["pp"], v["business"], v["net"], v["occupancy"]])
    total, last = _portfolio_values(conn, s, e), _portfolio_values(conn, ls, le)
    monthly.append([str(year), total["revenue"], total["operated"], total["pp"], total["business"], total["net"], total["occupancy"]])
    summary = {"heading": "Year summary", "columns": [_col("Item"), _col("This year", "mixed"), _col("Last year", "mixed"), _col("Change", "delta")],
               "rows": [], "row_fmts": [], "row_styles": [], "note": PORTFOLIO_NOTE.replace(DELTA_NOTE, "").strip()}
    for label, key, fmt, min_base, style in PORTFOLIO_LINES + [("Occupancy", "occupancy", "pct", 5, "")]:
        summary["rows"].append([label, total[key], last[key], pct_delta(total[key], last[key], min_base=min_base) if fmt == "money" else None])
        summary["row_fmts"].append(fmt); summary["row_styles"].append(style)
    rep["hero"] = [{"label": "Company Net Profit", "value": total["net"], "fmt": "money"}, {"label": "Property Profit", "value": total["pp"], "fmt": "money"}]
    rep["sections"] = [x for x in (summary,
                       {"heading": "Month by month", "columns": [_col("Month"), _col("Company Revenue", "money"), _col("Operated property costs", "money"), _col("Property Profit", "money"),
                                                                  _col("Business Costs", "money"), _col("Company Net Profit", "money"), _col("Occupancy", "pct")], "rows": monthly,
                        "row_styles": [""] * (len(monthly) - 1) + ["total"]},
                       _cost_centres(conn, s, e), _by_property(conn, s, e, heading=f"By property \u2014 {year}")) if x]
    return rep


def build_expense(conn, from_ym, to_ym, property_id=None):
    fy, fm = map(int, from_ym.split("-"))
    ty, tm = map(int, to_ym.split("-"))
    s, e = kpis.range_bounds(fy, fm, ty, tm)
    prop = get_property(conn, property_id) if property_id else None
    scope = short_name(prop["name"]) if prop else "All properties"
    period = f"{MONTH_NAMES[fm]} {fy}" if (fy, fm) == (ty, tm) else f"{MONTH_NAMES[fm]} {fy} to {MONTH_NAMES[tm]} {ty}"
    fperiod = from_ym if from_ym == to_ym else f"{from_ym}_to_{to_ym}"
    rep = _base("expense", scope, period, conn, filename_period=fperiod)
    if not completeness.period_has_data(conn, property_id, s, e):
        return _state_report(rep, "no_data", "No data has been recorded for this selection.")
    opex, capex = kpis.costs(conn, property_id, s, e, capex=False), kpis.costs(conn, property_id, s, e, capex=True)
    totals = {"heading": "Totals", "columns": [_col("Measure"), _col("Amount", "money")],
              "rows": [["Total recorded expenses", opex + capex], ["Opex", opex], ["Capex", capex]], "row_styles": ["total", "", ""],
              "note": "Every recorded expense transaction in the selection: operated property costs, managed properties' owner-side costs, management fees and Business Costs. "
                      "The Expenses page separates these and excludes the management fee, which is the company's income."}
    rep["hero"] = [{"label": "Total recorded expenses", "value": opex + capex, "fmt": "money"}]
    rep["sections"] = [x for x in (totals, _cost_centres(conn, s, e, property_id, managed=is_managed(prop) if prop else None), _category_table(conn, s, e, property_id),
                                   _vendor_table(conn, s, e, property_id), _ledger_table(conn, s, e, property_id, heading="Transaction ledger")) if x]
    return rep


# ------------------------------------------------------------------ renderers

def rows_for_display(section):
    """Formatted string rows (per-row formats when the column is 'mixed')."""
    out = []
    row_fmts = section.get("row_fmts")
    for i, row in enumerate(section["rows"]):
        cells = []
        for col, val in zip(section["columns"], row):
            fmt = row_fmts[i] if col["fmt"] == "mixed" and row_fmts else col["fmt"]
            cells.append(format_cell(val, fmt))
        out.append(cells)
    return out


def hero_display(report):
    return [{**h, "text": format_cell(h["value"], h["fmt"])} for h in report.get("hero", [])]


def _headline_section(report):
    """The headline figures as a section, so CSV / Excel / PDF carry exactly what the preview shows at the top."""
    if not report.get("hero"):
        return None
    return {"heading": "Headline", "columns": [_col("Item"), _col("Amount", "money"), _col("Basis")],
            "rows": [[h["label"], h["value"], h.get("note") or ""] for h in report["hero"]]}


def _export_sections(report):
    head = _headline_section(report)
    return ([head] if head else []) + report["sections"]


def _raw(v):
    return "" if v is None else (round(v, 2) if isinstance(v, float) else v)


def to_csv(report):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([report["title"], report["subtitle"]])
    if report.get("state"):
        w.writerow([]); w.writerow([report["state"]["text"]])
    for sec in _export_sections(report):
        w.writerow([])
        w.writerow([sec["heading"]])
        w.writerow([c["label"] for c in sec["columns"]])
        for row in sec["rows"]:
            w.writerow([_raw(v) for v in row])
        if sec.get("note"):
            w.writerow([sec["note"]])
    w.writerow([])
    for k, v in report.get("footer", []):
        w.writerow([k, v])
    return buf.getvalue()


def to_xlsx(report):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sections = _export_sections(report) or [{"heading": "Report", "columns": [_col("Item")], "rows": []}]
    for i, sec in enumerate(sections, 1):
        name = "".join(ch for ch in sec["heading"] if ch not in '[]:*?/\\')[:28] or f"Sheet{i}"
        ws = wb.create_sheet(f"{i}. {name}"[:31])
        ws.append([report["title"], report["subtitle"]])
        ws["A1"].font = Font(bold=True)
        ws.append([sec["heading"]])
        ws["A2"].font = Font(bold=True)
        ws.append([c["label"] for c in sec["columns"]])
        for cell in ws[3]:
            cell.font = Font(bold=True)
        for row in sec["rows"]:
            ws.append([_raw(v) for v in row])
        if sec.get("note"):
            ws.append([]); ws.append([sec["note"]])
        if i == len(sections):
            ws.append([])
            for k, v in report.get("footer", []):
                ws.append([k, v])
            if report.get("state"):
                ws.append([report["state"]["text"]])
        for col_cells in ws.columns:
            ws.column_dimensions[col_cells[0].column_letter].width = min(
                max(len(str(c.value)) if c.value is not None else 0 for c in col_cells[:60]) + 2, 48)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(report):
    from xml.sax.saxutils import escape as _escape

    def escape(text):
        """The PDF's standard fonts have no true minus sign: it is written as a hyphen-minus there (the amount and its sign are unchanged)."""
        return _escape(str(text).replace(MINUS, "-"))

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    ink, mute, line = colors.HexColor("#14201c"), colors.HexColor("#6b7872"), colors.HexColor("#d9e1dd")
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=8, textColor=mute, leading=11)
    kick = ParagraphStyle("kick", parent=styles["Normal"], fontSize=8, textColor=mute, leading=11)
    title = ParagraphStyle("t", parent=styles["Title"], alignment=0, fontSize=20, leading=24, textColor=ink, spaceAfter=2)
    head = ParagraphStyle("h", parent=styles["Heading3"], fontSize=11, textColor=ink, spaceBefore=6, spaceAfter=3)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
                            title=report["title"] + " \u2014 " + report["subtitle"])
    story = [Paragraph(escape(f"FERNHILL STAYS \u00b7 {report['title'].upper()}"), kick), Paragraph(escape(report["scope"]), title),
             Paragraph(escape(report["period"] + (f" \u00b7 {report['model'].capitalize()}" if report.get("model") else "")), styles["Normal"]), Spacer(1, 4 * mm)]
    if report.get("hero"):
        cells = [[Paragraph(escape(h["label"]) + (f" <font color='#6b7872'>({escape(h['note'])})</font>" if h.get("note") else ""), small) for h in report["hero"]],
                 [Paragraph(f"<font size=18><b>{escape(h['text'])}</b></font>", styles["Normal"]) for h in hero_display(report)]]
        t = Table(cells, hAlign="LEFT", colWidths=[70 * mm] * len(report["hero"]))
        t.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.5, line), ("LINEBELOW", (0, -1), (-1, -1), 0.5, line), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story += [t, Spacer(1, 4 * mm)]
    if report.get("state"):
        story.append(Paragraph(escape(report["state"]["text"]), styles["Italic"]))
    for sec in report["sections"]:
        story.append(Paragraph(escape(sec["heading"]), head))
        display = rows_for_display(sec)
        if not display:
            story += [Paragraph("No data for this period.", styles["Italic"]), Spacer(1, 3 * mm)]
            continue
        cellp = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8, leading=10)
        cellr = ParagraphStyle("cellr", parent=cellp, alignment=2)
        headp = ParagraphStyle("headp", parent=cellp, fontName="Helvetica-Bold", textColor=mute)
        headr = ParagraphStyle("headr", parent=headp, alignment=2)
        cols = sec["columns"]
        data = [[Paragraph(escape(c["label"]), headp if c["fmt"] == "text" else headr) for c in cols]] + [
            [Paragraph(("<b>%s</b>" if (sec.get("row_styles") or [""] * len(display))[ri] in ("total", "sub") else "%s") % escape(cell), cellp if cols[i]["fmt"] == "text" else cellr)
             for i, cell in enumerate(row)] for ri, row in enumerate(display)]
        weights = [3.0 if i == 0 else (2.0 if c["fmt"] == "text" else 1.25) for i, c in enumerate(cols)]
        widths = [174 * mm * w / sum(weights) for w in weights]
        table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        style = [("LINEBELOW", (0, 0), (-1, 0), 0.5, ink), ("LINEBELOW", (0, 1), (-1, -1), 0.25, line), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
        for ri, st in enumerate(sec.get("row_styles") or [], 1):
            if st == "total":
                style.append(("LINEABOVE", (0, ri), (-1, ri), 0.5, ink))
        table.setStyle(TableStyle(style))
        story.append(table)
        if sec.get("note"):
            story.append(Paragraph(escape(sec["note"]), small))
        story.append(Spacer(1, 4 * mm))
    foot = " \u00b7 ".join(f"{k}: {v}" for k, v in report.get("footer", []))
    story += [Spacer(1, 2 * mm), Paragraph(escape(foot), small)]
    doc.build(story)
    return buf.getvalue()
